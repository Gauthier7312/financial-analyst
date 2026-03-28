from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()

DARK  = "1A3A5C"
GREEN = "1E8449"
RED   = "C0392B"
GOLD  = "F1C40F"
LGREY = "F2F3F4"
WHITE = "FFFFFF"
LGRN  = "D5F5E3"
LRED  = "FADBD8"
LBLU  = "D6EAF8"

def tb():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def hf(sz=10, bold=True, color=WHITE):
    return Font(name="Arial", size=sz, bold=bold, color=color)

def bf(sz=10, bold=False, color="000000"):
    return Font(name="Arial", size=sz, bold=bold, color=color)

def fl(c): return PatternFill("solid", fgColor=c)
def ca(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
def la(): return Alignment(horizontal="left",   vertical="center", wrap_text=True)

def hrow(ws, row, cols, bg=DARK):
    for c in cols:
        cell = ws.cell(row=row, column=c)
        cell.font = hf(); cell.fill = fl(bg); cell.alignment = ca(); cell.border = tb()

def title(ws, rng, text, bg=DARK, sz=13):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = text; c.font = hf(sz=sz); c.fill = fl(bg); c.alignment = ca()

def wcol(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v


# ══════════════════════════════════════════════
# ONGLET 1 — JOURNAL DES TRADES
# ══════════════════════════════════════════════
ws1 = wb.active
ws1.title = "📒 Journal Trades"
ws1.row_dimensions[1].height = 38

title(ws1, "A1:O1", "JOURNAL DE BORD DES TRADES — BRVM 2026", sz=13)

heads = ["#", "Date", "Valeur", "Ticker", "Action", "Qté", "Prix unitaire",
         "Montant total", "Courtage (0.8%)", "Montant net", "Stop-loss",
         "Objectif CT", "Objectif MT", "Statut", "Notes"]
for c, h in enumerate(heads, 1):
    ws1.cell(row=2, column=c).value = h
hrow(ws1, 2, range(1, 16))

exemples = [
    (1, "À remplir", "SONATEL",  "SNTS.sn", "ACHAT", 7,  28400, "=F3*G3", "=H3*0.008", "=H3+I3", 24140, 30672, 32660, "⏳ Ouvert", "Tranche 1 DCA"),
    (2, "À remplir", "BOA CI",   "BOAC.ci", "ACHAT", 20,  8700, "=F4*G4", "=H4*0.008", "=H4+I4",  7395,  9744, 10875, "⏳ Ouvert", "Tranche 1 DCA"),
    (3, "À remplir", "TOTAL CI", "TTLC.ci", "ACHAT", 42,  2975, "=F5*G5", "=H5*0.008", "=H5+I5",  2529,  3272,  3570, "⏳ Ouvert", "Tranche 2 DCA"),
]

for i, row in enumerate(exemples):
    r = i + 3
    ws1.row_dimensions[r].height = 22
    bg = LGRN if row[4] == "ACHAT" else LRED
    for c, val in enumerate(row, 1):
        cell = ws1.cell(row=r, column=c)
        cell.value = val; cell.fill = fl(bg); cell.border = tb(); cell.alignment = ca()
        cell.font = bf(bold=(c in [3, 5, 14]))
        if c == 5:
            cell.font = Font(name="Arial", size=10, bold=True,
                             color=GREEN if val == "ACHAT" else RED)

for r in range(6, 56):
    ws1.row_dimensions[r].height = 20
    for c in range(1, 16):
        cell = ws1.cell(row=r, column=c)
        cell.fill = fl(LGREY if r % 2 == 0 else WHITE)
        cell.border = tb(); cell.alignment = ca(); cell.font = bf()

wcol(ws1, {"A":5,"B":14,"C":16,"D":12,"E":9,"F":7,"G":16,"H":16,
           "I":16,"J":16,"K":14,"L":14,"M":14,"N":12,"O":28})


# ══════════════════════════════════════════════
# ONGLET 2 — DIVIDENDES REÇUS
# ══════════════════════════════════════════════
ws2 = wb.create_sheet("💰 Dividendes")
ws2.row_dimensions[1].height = 35
title(ws2, "A1:H1", "SUIVI DES DIVIDENDES REÇUS — 2026", bg=GREEN)

heads2 = ["Date détachement", "Date paiement", "Valeur", "Nb actions",
          "Dividende/action", "Montant brut", "Impôt (si applicable)", "Montant net reçu"]
for c, h in enumerate(heads2, 1):
    ws2.cell(row=2, column=c).value = h
hrow(ws2, 2, range(1, 9), bg=GREEN)

divs = [
    ("21/04/2026", "~05/05/2026", "BOA BF",   "À préciser", 397,  "",  0, ""),
    ("22/05/2026", "~05/06/2026", "SONATEL",  "7",          1740, "=D4*E4", 0, "=F4-G4"),
    ("28/05/2026", "~10/06/2026", "BOA SN",   "À préciser", 450,  "",  0, ""),
    ("À préciser", "À préciser",  "ORANGE CI","À préciser", 704,  "",  0, ""),
    ("À préciser", "À préciser",  "PALMCI",   "À préciser", 442,  "",  0, ""),
    ("À préciser", "À préciser",  "SAPH CI",  "À préciser", 430,  "",  0, ""),
]

for i, row in enumerate(divs):
    r = i + 3
    ws2.row_dimensions[r].height = 22
    bg = LGRN if i % 2 == 0 else WHITE
    for c, val in enumerate(row, 1):
        cell = ws2.cell(row=r, column=c)
        cell.value = val; cell.fill = fl(bg); cell.border = tb()
        cell.alignment = ca(); cell.font = bf(bold=(c == 3))

# Total
rt = len(divs) + 4
ws2.cell(row=rt, column=1).value = "TOTAL DIVIDENDES ATTENDUS 2026"
ws2.merge_cells(f"A{rt}:E{rt}")
ws2.cell(row=rt, column=1).font = hf(color=WHITE)
ws2.cell(row=rt, column=1).fill = fl(GREEN)
ws2.cell(row=rt, column=1).alignment = ca()
ws2.cell(row=rt, column=8).value = f"=SUM(H3:H{rt-1})"
ws2.cell(row=rt, column=8).font = hf(color=WHITE)
ws2.cell(row=rt, column=8).fill = fl(GREEN)
ws2.cell(row=rt, column=8).alignment = ca()

wcol(ws2, {"A":18,"B":16,"C":14,"D":14,"E":18,"F":16,"G":18,"H":18})


# ══════════════════════════════════════════════
# ONGLET 3 — PERFORMANCE MENSUELLE
# ══════════════════════════════════════════════
ws3 = wb.create_sheet("📈 Performance Mensuelle")
ws3.row_dimensions[1].height = 35
title(ws3, "A1:J1", "SUIVI DE PERFORMANCE MENSUELLE — 2026/2027")

heads3 = ["Mois", "BRVM Composite", "Var. Indice", "Valeur Portef.",
          "Var. Portef.", "Dividendes", "Portef. Total", "vs Indice (alpha)", "Trades du mois", "Notes"]
for c, h in enumerate(heads3, 1):
    ws3.cell(row=2, column=c).value = h
hrow(ws3, 2, range(1, 11))

mois = ["Mars 2026","Avr. 2026","Mai 2026","Juin 2026","Juil. 2026",
        "Août 2026","Sep. 2026","Oct. 2026","Nov. 2026","Déc. 2026",
        "Jan. 2027","Fév. 2027","Mars 2027"]
valeurs_depart = [406.18, "", "", "", "", "", "", "", "", "", "", "", ""]
portef_depart  = [500000, "", "", "", "", "", "", "", "", "", "", "", ""]

for i, (m, vi, pi) in enumerate(zip(mois, valeurs_depart, portef_depart)):
    r = i + 3
    ws3.row_dimensions[r].height = 22
    bg = LGREY if i % 2 == 0 else WHITE
    row_data = [m, vi, "", pi, "", "", "", "", "", ""]
    for c, val in enumerate(row_data, 1):
        cell = ws3.cell(row=r, column=c)
        cell.value = val; cell.fill = fl(bg); cell.border = tb()
        cell.alignment = ca(); cell.font = bf(bold=(c == 1))
    if i > 0:
        ws3.cell(row=r, column=3).value = f"=SI(B{r}<>0,(B{r}-B{r-1})/B{r-1},\"\")"
        ws3.cell(row=r, column=5).value = f"=SI(D{r}<>0,(D{r}-D{r-1})/D{r-1},\"\")"
        ws3.cell(row=r, column=7).value = f"=SI(D{r}<>\"\",D{r}+F{r},\"\")"
        ws3.cell(row=r, column=8).value = f"=SI(E{r}<>\"\",E{r}-C{r},\"\")"

wcol(ws3, {"A":14,"B":18,"C":14,"D":16,"E":14,"F":14,"G":16,"H":18,"I":16,"J":26})


# ══════════════════════════════════════════════
# ONGLET 4 — TABLEAU DE BORD
# ══════════════════════════════════════════════
ws4 = wb.create_sheet("🎯 Tableau de Bord")
ws4.row_dimensions[1].height = 40
title(ws4, "A1:F1", "TABLEAU DE BORD — VUE D'ENSEMBLE", sz=14)

title(ws4, "A3:F3", "PORTEFEUILLE EN TEMPS RÉEL", bg="1E8449", sz=11)

p_heads = ["Valeur", "Allocation", "Investi", "Cours actuel", "Valeur actuelle", "Gain/Perte %"]
for c, h in enumerate(p_heads, 1):
    ws4.cell(row=4, column=c).value = h
hrow(ws4, 4, range(1, 7), bg="1E8449")

positions = [("SONATEL", "40%", 200000, 28400), ("BOA CI", "35%", 175000, 8700), ("TOTAL CI", "25%", 125000, 2975)]
for i, (name, alloc, inv, cours) in enumerate(positions):
    r = i + 5
    bg = [LGRN, LBLU, "FDEBD0"][i]
    ws4.cell(row=r, column=1).value = name
    ws4.cell(row=r, column=2).value = alloc
    ws4.cell(row=r, column=3).value = inv
    ws4.cell(row=r, column=4).value = cours
    ws4.cell(row=r, column=5).value = f"À mettre à jour"
    ws4.cell(row=r, column=6).value = f"À calculer"
    ws4.row_dimensions[r].height = 24
    for c in range(1, 7):
        ws4.cell(row=r, column=c).fill = fl(bg)
        ws4.cell(row=r, column=c).border = tb()
        ws4.cell(row=r, column=c).alignment = ca()
        ws4.cell(row=r, column=c).font = bf(bold=(c == 1))

# KPIs
r_kpi = 10
title(ws4, f"A{r_kpi}:F{r_kpi}", "KPIs CLÉS", bg=DARK, sz=11)
kpis = [
    ("Capital total investi", "500 000 FCFA", "Hors frais de courtage"),
    ("Courtage payé total", "~4 000 FCFA", "0.8% × 500K (estimation)"),
    ("Conservation annuelle", "~1 250 FCFA", "0.25% × 500K"),
    ("Tenue de compte annuelle", "10 000 FCFA", "2 500 × 4 trimestres"),
    ("Coût total annuel", "~15 250 FCFA", "= 3.05% du capital"),
    ("Rendement min. à battre", "> 3.05%", "Pour être profitable net de frais"),
    ("Dividendes attendus 2026", "~12 000+ FCFA", "SONATEL seul = 12 180 FCFA (7 actions)"),
    ("Objectif rendement annuel", "9–14%", "Base réaliste sur BRVM bull market"),
]

for i, (k, v, note) in enumerate(kpis):
    r = r_kpi + 1 + i
    ws4.row_dimensions[r].height = 22
    ws4.cell(row=r, column=1).value = k
    ws4.merge_cells(f"A{r}:A{r}")
    ws4.cell(row=r, column=2).value = v
    ws4.merge_cells(f"C{r}:F{r}")
    ws4.cell(row=r, column=3).value = note
    bg = LGREY if i % 2 == 0 else WHITE
    for c in [1, 2]:
        ws4.cell(row=r, column=c).fill = fl(bg)
        ws4.cell(row=r, column=c).border = tb()
        ws4.cell(row=r, column=c).alignment = ca()
        ws4.cell(row=r, column=c).font = bf(bold=(c == 1), color=DARK if c == 1 else "0000FF")
    ws4.cell(row=r, column=3).fill = fl(bg)
    ws4.cell(row=r, column=3).border = tb()
    ws4.cell(row=r, column=3).alignment = la()
    ws4.cell(row=r, column=3).font = bf(color="555555")

wcol(ws4, {"A":28,"B":22,"C":18,"D":16,"E":18,"F":16})

# Save
path = "/Users/elyseebleu/Documents/financial-analyst/journal_trades.xlsx"
wb.save(path)
print(f"✅ journal_trades.xlsx créé : {path}")
